# -*- coding: utf-8 -*-
"""
Ekspor RPS dengan mengisi LANGSUNG file template Word asli (assets/rps_template.docx),
bukan membangun ulang tabel dari kode. Prinsipnya: hanya mengganti teks di sel-sel yang
memang perlu diisi; struktur, warna, font, border sama sekali tidak disentuh karena
memang tidak pernah dibuat ulang.

PDF dihasilkan dengan mengonversi file Word yang sudah terisi lewat LibreOffice
(command-line `soffice`), supaya PDF dan Word dijamin identik satu sama lain.
"""

import io
import os
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy

from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from pdf_export import with_code, generate_qr_image, build_signature_qr_text

ASSETS_DIR = os.path.join(os.path.dirname(__file__), "assets")
TEMPLATE_PATH = os.path.join(ASSETS_DIR, "rps_template.docx")

# Dipindah ke sini (dari app.py) supaya bisa dipakai bersama oleh app.py sendiri
# dan rps_browse.py (halaman "RPS Disetujui") tanpa duplikasi/import silang.
BOBOT_KATEGORI = {"Interaksi": 30, "UTS": 20, "Tugas": 20, "UAS": 30}

# Urutan minggu -> indeks baris tabel pertemuan pada template BARU. Berbeda dari
# template lama: minggu ke-8 (UTS) sekarang PUNYA barisnya sendiri (tidak dilompati),
# dan minggu ke-16 (UAS) juga baris biasa (bukan lagi 1 baris gabungan bertuliskan
# "UAS" doang) - jadi pemetaannya jadi sesederhana m -> m+2 (baris 0-1 = header/spacer).
MINGGU_TO_ROW = {m: m + 2 for m in range(1, 17)}


def set_table_fixed_layout(table):
    """Paksa tabel memakai lebar kolom TETAP (bukan 'autofit' bawaan Word/LibreOffice).

    Tabel di file template asli tidak menetapkan w:tblLayout sama sekali, artinya
    defaultnya 'autofit': lebar kolom menyesuaikan panjang isi teks. Selama sel-selnya
    kosong (seperti saat diunggah), ini tidak kelihatan masalahnya - begitu diisi teks
    yang lebih panjang dari placeholder kosong, tabel jadi melebar dan sebagian kolom
    terpotong keluar dari tepi halaman. Fixed layout membuat lebar kolom tetap sesuai
    tblGrid asli, dan teks yang panjang akan membungkus ke baris berikutnya (wrap)
    alih-alih melebarkan tabel.
    """
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_paragraph_text(p, text):
    """Ganti isi SATU paragraf (bukan seluruh sel) dengan teks baru, sambil
    mempertahankan format 'mark' bawaan placeholder-nya. Dipisah dari
    set_cell_text() supaya bisa dipakai untuk sel QR di template baru yang
    punya 2 paragraf terpisah dalam satu sel (gambar QR di paragraf 1, nama
    di paragraf 2) tanpa saling menghapus."""
    text = "" if text is None else str(text)

    # Sel placeholder kosong di template biasanya sudah punya "format bawaan" (font,
    # ukuran, bold, dst) tersimpan di w:pPr/w:rPr paragrafnya - ini yang menentukan
    # tampilan teks BARU yang diketik di situ (di Word, ini keliatan halus: taruh
    # kursor di sel kosong, formatnya sudah "siap pakai" walau belum ada teks apa
    # pun). Run baru yang dibuat lewat python-docx TIDAK otomatis mewarisi ini -
    # kalau tidak disalin manual, hasilnya balik ke format default Word (mis. ukuran
    # font beda dari yang seharusnya). Disalin di sini supaya konsisten.
    mark_rpr = None
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None:
        rpr_el = pPr.find(qn("w:rPr"))
        if rpr_el is not None:
            mark_rpr = deepcopy(rpr_el)

    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    # Beberapa sel contoh di template asli (mis. baris minggu 1) sudah punya format
    # daftar bernomor otomatis (numPr) bawaan. Kalau tidak dibuang, teks baru yang
    # sudah kita tulis manual dengan "1." "2." akan tertumpuk jadi "1.1." "1.2." dst.
    if pPr is not None:
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            pPr.remove(numPr)
    lines = text.split("\n")
    run = p.add_run(lines[0])
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)

    if mark_rpr is not None:
        # w:rPr salinan ini asalnya dari w:pPr (properti "mark" paragraf), pindahkan
        # jadi w:rPr milik run yang baru dibuat supaya font/ukuran/bold-nya kepakai.
        r_el = run._element
        existing_rpr = r_el.find(qn("w:rPr"))
        if existing_rpr is not None:
            r_el.remove(existing_rpr)
        r_el.insert(0, mark_rpr)


def set_cell_text(cell, text):
    """Ganti isi sel TANPA mengubah struktur tabel - hanya teks di paragraf pertama
    (paragraf LAIN di sel yang sama ikut dihapus). Untuk sel dengan >1 paragraf yang
    PERLU dipertahankan (mis. sel QR di template baru), pakai set_paragraph_text()
    langsung ke paragraf yang dimaksud, bukan fungsi ini."""
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    set_paragraph_text(cell.paragraphs[0], text)


def dedup_row_cells(row):
    """Sel unik per baris (menghapus duplikat akibat gridSpan horizontal)."""
    result = []
    last_path = None
    for c in row.cells:
        path = c._tc.getroottree().getpath(c._tc)
        if path != last_path:
            result.append(c)
            last_path = path
    return result


def fill_qr_cell(cell, aksi, kode_mk, nama_mk, nama_orang, peran, tanggal):
    """Isi satu sel QR bawaan TEMPLATE BARU - selnya sudah punya 2 paragraf sendiri
    ([0]=placeholder gambar QR, [1]=placeholder nama), diverifikasi lewat XML asli.
    ini MENGGANTIKAN blok tanda tangan+QR terpisah di akhir dokumen dari template
    lama (dulu ditambahkan programatik lewat append_signature_block karena template
    resmi belum punya tempatnya sendiri - sekarang sudah disediakan template, jadi
    fungsi append tsb sudah dihapus, tidak dipakai lagi)."""
    paras = cell.paragraphs
    qr_p = paras[0]
    for r in list(qr_p.runs):
        r._element.getparent().remove(r._element)
    qr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr_text = build_signature_qr_text(aksi, kode_mk, nama_mk, nama_orang, peran, tanggal)
    qr_p.add_run().add_picture(generate_qr_image(qr_text), width=Cm(2.0))
    if len(paras) > 1:
        set_paragraph_text(paras[1], nama_orang or "(belum diisi)")


def fill_info_table(document, prodi, mk_row, cpl_df, info_umum, cpmk_data,
                     komponen_data, bobot_kategori, referensi_data):
    t = document.tables[0]

    # --- judul (baris 0) & tanggal penyusunan (baris 1) ---
    title_cell = dedup_row_cells(t.rows[0])[-1]
    original_title = title_cell.text
    if "….." in original_title:
        full_title = original_title.replace("…..", prodi)
    elif "....." in original_title:
        full_title = original_title.replace(".....", prodi)
    else:
        parts = original_title.split("\n")
        if len(parts) >= 2:
            parts[1] = f"PROGRAM STUDI {prodi}"
        full_title = "\n".join(parts)
    set_cell_text(title_cell, full_title)

    tgl_cell = dedup_row_cells(t.rows[1])[0]
    tgl = info_umum.get("tanggal_dokumen", "") or "-"
    set_cell_text(tgl_cell, f"Tanggal Penyusunan : {tgl}")

    # --- identitas (baris 2-8); pemetaan sel sudah diverifikasi lewat XML template
    #     baru. Baris "MK yang menjadi prasyarat"/"Menjadi Prasyarat untuk MK" dari
    #     template lama sudah TIDAK ADA lagi di template baru - tidak diisi lagi. ---
    r2 = t.rows[2].cells
    set_cell_text(r2[1], mk_row.get("Nama Mata Kuliah", "-"))

    r3 = t.rows[3].cells
    set_cell_text(r3[1], str(mk_row.get("Kode MK", "-")))

    r4 = t.rows[4].cells
    set_cell_text(r4[1], info_umum.get("rumpun_mk", "") or "-")

    r5 = t.rows[5].cells
    set_cell_text(r5[1], f"{mk_row.get('SKS', '-')} SKS")

    r6 = t.rows[6].cells
    set_cell_text(r6[1], str(mk_row.get("Semester", "-")))

    r7 = t.rows[7].cells
    # "Dosen Pengampu" SENGAJA diambil dari info_umum apa adanya di sini (bukan
    # dihitung ulang) - pemanggil (app.py/rps_browse.py) yang bertanggung jawab
    # mengisi nilai ini dengan nama akun yang SEDANG mengunduh dokumen, sesuai
    # keputusan: "Dosen Pengampu" di dokumen = siapa pun yang mengunduhnya saat itu,
    # bukan disimpan permanen di data RPS.
    set_cell_text(r7[1], info_umum.get("dosen_pengampu", "") or "-")

    r8 = t.rows[8].cells
    set_cell_text(r8[1], info_umum.get("deskripsi_mk", "") or "-")

    # --- QR Koordinator, Kaprodi & BPM (baris 3, kolom 3/5/7) - menggantikan
    #     blok tanda tangan terpisah di akhir dokumen dari template lama.
    #     PENTING: index kolom BERGESER dari revisi template sebelumnya karena
    #     tabel identitas sekarang 8 kolom (dulu 7) untuk menampung slot BPM -
    #     koordinator 4->3, kaprodi 6->5 (diverifikasi ulang lewat XML mentah,
    #     bukan tebakan). Selnya di-vertical-merge menutupi baris 3-7 (Kode s/d
    #     Dosen Pengampu), tapi konten sungguhan cuma ada di baris asal (baris
    #     3) - baris 4-7 cuma penanda lanjutan merge, TIDAK PERLU ditulis lagi. ---
    kode_mk = str(mk_row.get("Kode MK", "-"))
    nama_mk = mk_row.get("Nama Mata Kuliah", "-")
    fill_qr_cell(
        r3[3], "disusun", kode_mk, nama_mk,
        info_umum.get("dosen_koordinator"), "Koordinator Mata Kuliah", tgl,
    )
    fill_qr_cell(
        r3[5], "disetujui", kode_mk, nama_mk,
        info_umum.get("nama_kaprodi"), "Ketua Program Studi", tgl,
    )
    fill_qr_cell(
        r3[7], "divalidasi", kode_mk, nama_mk,
        info_umum.get("nama_biro_pjm"), "Ka. Biro Penjaminan Mutu", tgl,
    )

    # --- CPL (baris 11-15) ---
    cpl_selected = [cpmk_data[i]["cpl_kode"] for i in range(1, 6)]
    for i, kode in enumerate(cpl_selected):
        row_cells = t.rows[11 + i].cells
        desk = cpl_df.loc[cpl_df["Kode CPL"] == kode, "Deskripsi CPL"]
        teks = desk.values[0] if len(desk) else ""
        set_cell_text(row_cells[0], kode)
        set_cell_text(row_cells[1], teks)

    # --- CPMK (baris 17-21) ---
    for i in range(1, 6):
        c = cpmk_data[i]
        row_cells = t.rows[16 + i].cells
        set_cell_text(row_cells[1], with_code(c["deskripsi"], c["cpl_kode"]))

    # --- Mapping CPMK -> Komponen Penilaian (baris 24-28); urutan kolom template:
    #     Interaksi | UTS | Tugas | UAS (kategori "Kehadiran dan Sikap" lama
    #     berganti nama jadi "Interaksi" - lihat BOBOT_KATEGORI). Kolom UTS,
    #     Tugas & UAS di file asli memakai gridSpan 2 kolom mentah, jadi WAJIB
    #     pakai sel yang sudah di-dedup. Tiap sel diisi PERSENTASE kontribusi
    #     CPMK tsb ke komponen itu (komponen_data[i] = {kategori: persen}).
    #     Total persentase satu kolom (dijumlah semua CPMK) seharusnya PAS sama
    #     dengan bobot_kategori kolom itu - divalidasi di app.py sebelum RPS
    #     diajukan, BUKAN di sini (fungsi ini murni menulis apa adanya). ---
    kolom_urutan = ["Interaksi", "UTS", "Tugas", "UAS"]
    for i in range(1, 6):
        row_cells = dedup_row_cells(t.rows[23 + i])
        cpmk_kat = komponen_data.get(i)
        if not isinstance(cpmk_kat, dict):
            # Kompatibilitas RPS lama (format checklist/nilai tunggal sebelum
            # jadi persentase) - tidak membawa info persentase sama sekali,
            # dianggap kosong (0 di semua kategori untuk CPMK ini).
            cpmk_kat = {}
        for j, kat in enumerate(kolom_urutan, start=1):
            nilai = cpmk_kat.get(kat) or 0
            set_cell_text(row_cells[j], f"{nilai}%" if nilai else "")

    # --- Bobot Nilai (baris 29) - total per kategori (mis. UTS 20%), BUKAN
    #     dijumlah dari komponen_data - ini murni angka baku bobot_kategori. ---
    bobot_row_cells = dedup_row_cells(t.rows[29])
    for j, kat in enumerate(kolom_urutan, start=1):
        set_cell_text(bobot_row_cells[j], f"{bobot_kategori.get(kat, 0)}%")

    # --- Level Penggunaan AI (baris 31) - BARU, template lama tidak punya baris
    #     ini. Selnya punya 2 paragraf terpisah (sama pola dengan sel QR):
    #     [0]=pilihan level, [1]=deskripsi. Opsional - kalau kosong, tampilkan
    #     placeholder yang jelas, bukan dibiarkan blank membingungkan. ---
    ai_cell = dedup_row_cells(t.rows[31])[1]
    ai_paras = ai_cell.paragraphs
    set_paragraph_text(ai_paras[0], info_umum.get("level_ai") or "(belum dipilih)")
    if len(ai_paras) > 1:
        set_paragraph_text(ai_paras[1], info_umum.get("deskripsi_ai") or "(belum diisi)")

    # --- Daftar Pustaka (baris 32 - bergeser lagi 1 baris karena ada tambahan
    #     baris Level Penggunaan AI tepat di atasnya) ---
    pustaka_lines = [ref["sitasi"] for ref in referensi_data if ref.get("sitasi")]
    pustaka_text = "\n".join(f"{i}. {t_}" for i, t_ in enumerate(pustaka_lines, start=1)) or "(belum diisi)"
    set_cell_text(t.rows[32].cells[1], pustaka_text)


def fill_pertemuan_table(document, pertemuan_data, sks):
    t = document.tables[1]
    # Keterangan Waktu = SKS x rasio menit/minggu baku (Tatap Muka 50,
    # Belajar Terstruktur 60, Belajar Mandiri 60 - sesuai contoh isian yang
    # sudah disiapkan di template: "[TM: 1x (SKSx50 menit)]" dst.). SAMA untuk
    # semua 16 minggu karena SKS satu mata kuliah memang tetap, bukan berubah
    # per minggu.
    try:
        sks_int = int(sks)
    except (TypeError, ValueError):
        sks_int = None
    if sks_int is not None:
        keterangan_waktu = (
            f"\n\nWaktu:\n"
            f"[TM: 1x ({sks_int}x50 menit)]\n"
            f"[BT: 1x ({sks_int}x60 menit)]\n"
            f"[BM: 1x ({sks_int}x60 menit)]"
        )
    else:
        keterangan_waktu = ""

    for minggu, row_idx in MINGGU_TO_ROW.items():
        p = pertemuan_data.get(minggu, {})
        cells = t.rows[row_idx].cells
        # urutan kolom BERSIH pada template baru (7 kolom, tanpa duplikasi gridSpan
        # seperti template lama): 0=Minggu 1=Sub-CPMK 2=Bloom 3=Indikator
        # 4=Bentuk Asesmen 5=Metode+Daring(gabungan, + Keterangan Waktu) 6=Materi
        sub_text = with_code(p.get("sub_cpmk_desc", ""), p.get("cpmk_ref"))
        set_cell_text(cells[1], sub_text)
        set_cell_text(cells[2], ", ".join(p.get("bloom", [])))
        set_cell_text(cells[3], p.get("indikator", ""))
        set_cell_text(cells[4], p.get("bentuk_asesmen", ""))
        metode_daring = ", ".join(list(p.get("metode", [])) + list(p.get("bentuk", [])))
        set_cell_text(cells[5], metode_daring + keterangan_waktu)
        set_cell_text(cells[6], p.get("materi", ""))



def apply_fixed_layout_everywhere(document):
    """Terapkan fixed layout ke semua tabel, termasuk yang bersarang di dalam sel
    (mis. tabel Pengertian 1 SKS, Metode Pembelajaran, Bloom Taxonomy)."""
    def recurse(table):
        set_table_fixed_layout(table)
        for row in table.rows:
            for cell in row.cells:
                for nested in cell.tables:
                    recurse(nested)
    for t in document.tables:
        recurse(t)


def build_docx(prodi, mk_row, cpl_df, info_umum, cpmk_data,
                pertemuan_data, referensi_data, komponen_data, bobot_kategori):
    """Isi template asli - TIDAK membangun tabel dari nol. Mengembalikan BytesIO .docx."""
    document = Document(TEMPLATE_PATH)
    apply_fixed_layout_everywhere(document)
    fill_info_table(document, prodi, mk_row, cpl_df, info_umum, cpmk_data,
                     komponen_data, bobot_kategori, referensi_data)
    fill_pertemuan_table(document, pertemuan_data, mk_row.get("SKS"))

    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf


def find_soffice():
    """Cari executable LibreOffice/soffice di lokasi umum (Windows/Mac/Linux).

    Sengaja TIDAK hanya mengandalkan PATH sistem, karena installer LibreOffice di
    Windows sering tidak otomatis menambahkan dirinya ke PATH - jadi meski sudah
    ter-install dengan benar, perintah 'soffice' bisa saja tetap 'not recognized'
    di Command Prompt/PowerShell. Di sini dicoba beberapa cara sekaligus: override
    manual (config/soffice_path.txt), PATH, lokasi instalasi umum (termasuk drive
    selain C:), variabel lingkungan Windows, dan Windows Registry (App Paths, cara
    paling andal karena diisi otomatis oleh installer resmi LibreOffice).
    """
    # 0) Override manual - kalau semua deteksi otomatis di bawah gagal, isi saja
    #    lokasi soffice.exe secara manual di file ini (satu baris, path lengkap).
    override_path = os.path.join(os.path.dirname(__file__), "config", "soffice_path.txt")
    if os.path.isfile(override_path):
        try:
            with open(override_path, "r", encoding="utf-8") as f:
                manual_path = f.read().strip()
            if manual_path and os.path.isfile(manual_path):
                return manual_path
        except Exception:
            pass

    # 1) PATH sistem (kalau memang sudah diset / di Mac & Linux biasanya begini)
    for name in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found

    candidates = []
    # 2) Lokasi instalasi umum di Windows, di berbagai drive yang lazim dipakai
    program_dirs = []
    for var in ("PROGRAMFILES", "PROGRAMFILES(X86)", "PROGRAMW6432"):
        val = os.environ.get(var)
        if val:
            program_dirs.append(val)
    for drive in ("C:", "D:", "E:"):
        program_dirs.append(f"{drive}\\Program Files")
        program_dirs.append(f"{drive}\\Program Files (x86)")
    seen = set()
    for base in program_dirs:
        base = base.rstrip("\\")
        if base.lower() in seen:
            continue
        seen.add(base.lower())
        candidates.append(f"{base}\\LibreOffice\\program\\soffice.exe")

    # 3) Lokasi umum di Mac & Linux
    candidates += [
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice", "/usr/bin/libreoffice",
        "/usr/local/bin/soffice", "/opt/libreoffice/program/soffice",
        "/snap/bin/libreoffice",
    ]
    for c in candidates:
        if os.path.isfile(c):
            return c

    # 4) Windows Registry (App Paths) - diisi otomatis oleh installer resmi
    #    LibreOffice, jadi ini cara paling andal kalau lokasi instalasinya custom.
    if sys.platform == "win32":
        try:
            import winreg
            for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
                for subkey in (
                    r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\soffice.exe",
                    r"SOFTWARE\WOW6432Node\Microsoft\Windows\CurrentVersion\App Paths\soffice.exe",
                ):
                    try:
                        with winreg.OpenKey(hive, subkey) as key:
                            path, _ = winreg.QueryValueEx(key, "")
                            if path and os.path.isfile(path):
                                return path
                    except FileNotFoundError:
                        continue
        except Exception:
            pass

    return None


def build_pdf_via_libreoffice(docx_bytes_io):
    """Konversi .docx (BytesIO) -> .pdf (BytesIO) lewat LibreOffice, supaya PDF
    dijamin identik dengan Word-nya. Mengembalikan None kalau LibreOffice tidak
    ditemukan (aplikasi akan menampilkan pesan yang jelas ke pengguna)."""
    soffice = find_soffice()
    if not soffice:
        return None

    with tempfile.TemporaryDirectory() as tmp:
        docx_path = os.path.join(tmp, "rps.docx")
        with open(docx_path, "wb") as f:
            f.write(docx_bytes_io.getvalue())
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, docx_path],
            capture_output=True, timeout=90,
        )
        pdf_path = os.path.join(tmp, "rps.pdf")
        if result.returncode != 0 or not os.path.exists(pdf_path):
            return None
        with open(pdf_path, "rb") as f:
            data = f.read()
        return io.BytesIO(data)